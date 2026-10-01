"""
Regressionstests für die manuelle Terminverwaltung im Reiter "Routenplanung"
(app.py, logic/route_import.py) - deckt die 10 im Auftrag geforderten
Testfälle ab (Termine werden primär DIREKT IN DER WEBAPP angelegt, ein
Tabellenblatt "Routenplanung" in der Excel-Datei bleibt optional).

Verwendet für reine Routenlogik (Fall 2, 5, 6, 7, 8) direkt build_termin() +
logic/route_planner.py mit gestubbten geocode_fn/travel_fn (keine echten
Netzwerkzugriffe) und für app.py-spezifisches Verhalten (Fall 1, 3, 4, 9, 10)
Streamlits AppTest-Framework, analog zu tests/test_app_rendering.py.

Ausführen:
    python -m unittest discover tests
"""

import datetime as dt
import io
import os
import unittest

import openpyxl
import pandas as pd
import yaml

from logic.geocoding import GeocodeResult, haversine_km
from logic.route_import import build_termin, next_termin_id
from logic.route_planner import compute_route, filter_termine_by_date

from streamlit.testing.v1 import AppTest

with open("config.yaml", "r", encoding="utf-8") as f:
    CONFIG = yaml.safe_load(f)

ORTE = {
    "Start": (48.0, 9.0),
    "Ort A": (48.1, 9.0),
    "Ort B": (48.2, 9.0),
    "Ort C": (48.3, 9.0),
}


def geocode_fn(address, cache=None):
    if cache is not None and address in cache:
        return cache[address]
    coord = ORTE.get(address)
    result = GeocodeResult(lat=coord[0], lon=coord[1], display_name=address) if coord else None
    if cache is not None:
        cache[address] = result
    return result


def travel_fn_60kmh(coord_a, coord_b):
    km = haversine_km(coord_a[0], coord_a[1], coord_b[0], coord_b[1])
    return km / 60 * 3600, km * 1000


class TestFall2DreiUnternehmenAlsTermineAnlegen(unittest.TestCase):
    """Fall 2: 3 Unternehmen aus Sales Intelligence auswählen, je Termin
    anlegen -> Routenplanung funktioniert (alle 3 werden berechnet)."""

    def test_drei_manuell_angelegte_termine_ergeben_drei_stops(self):
        termine = [
            build_termin(
                "Hotel A", "Teststr. 1", "00000", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0),
                dirs21_id="1", termin_id=next_termin_id(),
            ),
            build_termin(
                "Hotel B", "Teststr. 1", "00000", "Ort B", dt.date(2026, 10, 15), dt.time(11, 0),
                dirs21_id="2", termin_id=next_termin_id(),
            ),
            build_termin(
                "Hotel C", "Teststr. 1", "00000", "Ort C", dt.date(2026, 10, 15), dt.time(13, 0),
                dirs21_id="3", termin_id=next_termin_id(),
            ),
        ]
        route = compute_route(
            termine, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
        )
        self.assertEqual(route["anzahl_termine"], 3)
        self.assertEqual([s["hotel_name"] for s in route["stops"]], ["Hotel A", "Hotel B", "Hotel C"])


class TestFall3DirsIdWirdUebernommen(unittest.TestCase):
    """Fall 3: bestehende DIRS21-ID -> automatisch übernommen (Auswahl eines
    Unternehmens im "Termin hinzufügen"-Formular befüllt Straße/PLZ/Ort über
    den on_change-Callback, die DIRS21-ID wird beim Absenden aus den
    Sales-Intelligence-Optionen übernommen, siehe app.py)."""

    def test_auswahl_befuellt_adresse_und_dirs21_id_bleibt_abrufbar(self):
        script = """
import streamlit as st
import app as app_module

st.session_state["analyzed_companies"] = [
    {"hotel_name": "Hotel Mit ID", "website": "mit-id.de", "strasse": "Hauptstr. 9", "plz": "88045"},
]
st.session_state["result_rows"] = [
    {"hotel_name": "Hotel Mit ID", "dirs21_id": "DIRS-42", "ort": "Friedrichshafen",
     "gesamtprioritaet": "A", "fachliche_top_empfehlung": "MICE"},
]

optionen = app_module._sales_company_options()
st.session_state["_route_form_optionen"] = optionen
st.session_state["route_form_unternehmen"] = "Hotel Mit ID"
app_module._on_route_form_company_change()

passendes = next(o for o in optionen if o["hotel_name"] == "Hotel Mit ID")
st.write(f"DIRS_ID={passendes['dirs21_id']}")
st.write(f"STRASSE={st.session_state.get('route_form_strasse')}")
st.write(f"PLZ={st.session_state.get('route_form_plz')}")
st.write(f"ORT={st.session_state.get('route_form_ort')}")
"""
        path = os.path.join(os.path.dirname(__file__), "_tmp_fall3.py")
        with open(path, "w", encoding="utf-8") as f:
            f.write(script)
        try:
            at = AppTest.from_file(path)
            at.run(timeout=30)
            self.assertEqual(len(at.exception), 0, msg=list(at.exception))
            texts = [m.value for m in at.markdown] + [t.value for t in at.text]
            self.assertIn("DIRS_ID=DIRS-42", texts)
            self.assertIn("STRASSE=Hauptstr. 9", texts)
            self.assertIn("PLZ=88045", texts)
            self.assertIn("ORT=Friedrichshafen", texts)
        finally:
            if os.path.exists(path):
                os.remove(path)


class TestFall4FehlendeStrasseManuellErgaenzbar(unittest.TestCase):
    """Fall 4: keine Straße vorhanden -> manuelle Ergänzung möglich, kein
    Absturz, Termin bleibt nutzbar (nur mit Prüfhinweis markiert)."""

    def test_termin_ohne_strasse_wird_markiert_route_trotzdem_berechnet(self):
        termine = [
            build_termin("Hotel Ohne Strasse", "", "00000", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0)),
        ]
        self.assertIn("Straße fehlt", termine[0]["pruefhinweis_import"])
        route = compute_route(
            termine, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
        )
        self.assertEqual(route["anzahl_termine"], 1)

        # manuelle Ergänzung über revalidate_termin (siehe app.py Bearbeiten-Formular)
        from logic.route_import import revalidate_termin
        termine[0]["strasse"] = "Nachträglich ergänzt"
        revalidate_termin(termine[0])
        self.assertEqual(termine[0]["pruefhinweis_import"], "")


class TestFall5FixerTerminBleibtUnveraendert(unittest.TestCase):
    """Fall 5: fixer Termin -> bleibt unverändert, auch bei der optimierten
    Reihenfolge."""

    def test_fixer_termin_behaelt_seine_uhrzeit_in_optimierter_route(self):
        termine = [
            build_termin("Hotel A", "T", "0", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0), termin_status="fix"),
            build_termin(
                "Hotel B", "T", "0", "Ort B", dt.date(2026, 10, 15), dt.time(10, 30), termin_status="flexibel",
                flexibel_von=dt.time(10, 0), flexibel_bis=dt.time(13, 0), termin_dauer_minuten=30,
            ),
            build_termin("Hotel C", "T", "0", "Ort C", dt.date(2026, 10, 15), dt.time(14, 0), termin_status="fix"),
        ]
        optimiert = compute_route(
            termine, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
            reihenfolge="optimiert",
        )
        fixe = [s for s in optimiert["stops"] if s["termin_status"] == "fix"]
        self.assertEqual({s["termin_beginn"] for s in fixe}, {dt.time(9, 0), dt.time(14, 0)})


class TestFall6FlexiblerTerminNurImZeitfenster(unittest.TestCase):
    """Fall 6: flexibler Termin -> wird nur innerhalb seines Zeitfensters
    verschoben, niemals davor/danach."""

    def test_flexibler_termin_landet_innerhalb_seines_fensters(self):
        termine = [
            build_termin("Hotel A", "T", "0", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0), termin_status="fix"),
            build_termin(
                "Hotel B", "T", "0", "Ort B", dt.date(2026, 10, 15), dt.time(9, 30), termin_status="flexibel",
                flexibel_von=dt.time(10, 0), flexibel_bis=dt.time(13, 0), termin_dauer_minuten=30,
            ),
            build_termin("Hotel C", "T", "0", "Ort C", dt.date(2026, 10, 15), dt.time(14, 0), termin_status="fix"),
        ]
        optimiert = compute_route(
            termine, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
            reihenfolge="optimiert",
        )
        flexibel = next(s for s in optimiert["stops"] if s["termin_status"] == "flexibel")
        beginn_min = flexibel["termin_beginn"].hour * 60 + flexibel["termin_beginn"].minute
        self.assertGreaterEqual(beginn_min, 10 * 60)
        self.assertLessEqual(beginn_min, 13 * 60)


class TestFall7TerminLoeschenRouteOhneIhn(unittest.TestCase):
    """Fall 7: Termin löschen -> Route wird danach ohne diesen Termin
    berechnet (Simulation: Löschen = Entfernen aus der Terminliste VOR dem
    erneuten Aufruf von compute_route, siehe app.py -> _render_terminliste
    "Löschen"-Button)."""

    def test_geloeschter_termin_erscheint_nicht_mehr_in_der_route(self):
        alle_termine = [
            build_termin("Hotel A", "T", "0", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0), termin_id="t1"),
            build_termin("Hotel B", "T", "0", "Ort B", dt.date(2026, 10, 15), dt.time(11, 0), termin_id="t2"),
        ]
        alle_termine = [t for t in alle_termine if t["termin_id"] != "t1"]
        termine_tag = filter_termine_by_date(alle_termine, dt.date(2026, 10, 15))
        route = compute_route(
            termine_tag, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
        )
        self.assertEqual(route["anzahl_termine"], 1)
        self.assertEqual(route["stops"][0]["hotel_name"], "Hotel B")


class TestFall8ExcelExportDreiGetrennteSheets(unittest.TestCase):
    """Fall 8: Excel-Download enthält Sales Intelligence, Routenplanung und
    Routen-Ergebnis als getrennte Tabellenblätter."""

    def test_workbook_enthaelt_alle_drei_sheets(self):
        import app as app_module
        from logic.exporter import export_workbook_bytes
        from logic.pipeline import RESULT_COLUMNS

        alle_termine = [
            build_termin("Hotel A", "Teststr. 1", "00000", "Ort A", dt.date(2026, 10, 15), dt.time(9, 0)),
        ]
        route = compute_route(
            alle_termine, "Start", "letzter_termin", None, CONFIG,
            geocode_fn=geocode_fn, travel_fn=travel_fn_60kmh, start_zeit=dt.time(8, 0), geocode_cache={},
        )
        sales_df = pd.DataFrame(columns=RESULT_COLUMNS)
        route_input_df = app_module._route_input_dataframe(alle_termine)
        route_ergebnis_df = app_module._route_ergebnis_dataframe(route)

        workbook_bytes = export_workbook_bytes(sales_df, route_input_df, route_ergebnis_df)
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_bytes))
        self.assertIn("DIRS21 Sales Intelligence", workbook.sheetnames)
        self.assertIn("Routenplanung", workbook.sheetnames)
        self.assertIn("Routen-Ergebnis", workbook.sheetnames)

        # Sales-Intelligence-Sheet wird NICHT um Routenfelder erweitert
        sales_header = [c.value for c in next(workbook["DIRS21 Sales Intelligence"].iter_rows(max_row=1))]
        self.assertNotIn("Termin_Datum", sales_header)


class TestFall1KeinRoutenplanungSheetStartetNormal(unittest.TestCase):
    """Fall 1 / Fall 10: Input-Excel ohne Tabellenblatt "Routenplanung" (bzw.
    gar keine hochgeladene Datei) -> die Webapp startet ohne Fehler, Termine
    können vollständig manuell angelegt werden."""

    def test_tab_ohne_hochgeladene_datei_zeigt_hinweis_kein_fehler(self):
        script = """
import app as app_module

config = app_module.load_config()
app_module.render_routenplanung_tab(config, None)
import streamlit as st
st.write("RENDER_COMPLETE")
"""
        path = os.path.join(os.path.dirname(__file__), "_tmp_fall1.py")
        with open(path, "w", encoding="utf-8") as f:
            f.write(script)
        try:
            at = AppTest.from_file(path)
            at.run(timeout=30)
            self.assertEqual(len(at.exception), 0, msg=list(at.exception))
            infos = [i.body for i in at.info]
            self.assertTrue(any("Noch keine Termine angelegt" in i for i in infos))
        finally:
            if os.path.exists(path):
                os.remove(path)


class TestFall9und10OptionalerSheetImport(unittest.TestCase):
    """Fall 9: Excel-Datei enthält bereits ein Tabellenblatt "Routenplanung"
    -> dessen Termine werden optional importiert. Fall 10: kein Tabellenblatt
    vorhanden -> kein Fehler, route_termine bleibt leer."""

    def _build_workbook_mit_sheet(self):
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            pd.DataFrame([{"Unternehmensname": "Dummy"}]).to_excel(writer, sheet_name="Sales Intelligence", index=False)
            pd.DataFrame([{
                "Unternehmensname": "Hotel Import", "DIRS21-ID": "9", "Straße": "Importstr. 1", "PLZ": "88045",
                "Ort": "Friedrichshafen", "Termin_Datum": "15.10.2026", "Termin_Uhrzeit": "10:00",
            }]).to_excel(writer, sheet_name="Routenplanung", index=False)
        buffer.seek(0)
        return buffer.getvalue()

    def _build_workbook_ohne_sheet(self):
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            pd.DataFrame([{"Unternehmensname": "Dummy"}]).to_excel(writer, sheet_name="Sales Intelligence", index=False)
        buffer.seek(0)
        return buffer.getvalue()

    def _run_tab_mit_bytes(self, uploaded_bytes):
        script_path = os.path.join(os.path.dirname(__file__), "_tmp_fall9_payload.xlsx")
        with open(script_path, "wb") as f:
            f.write(uploaded_bytes)
        script = f"""
import app as app_module
import streamlit as st

with open(r"{script_path}", "rb") as f:
    uploaded_bytes = f.read()

config = app_module.load_config()
app_module.render_routenplanung_tab(config, uploaded_bytes)
st.write(f"ANZAHL_TERMINE={{len(st.session_state.get('route_termine', []))}}")
st.write("RENDER_COMPLETE")
"""
        app_path = os.path.join(os.path.dirname(__file__), "_tmp_fall9.py")
        with open(app_path, "w", encoding="utf-8") as f:
            f.write(script)
        try:
            at = AppTest.from_file(app_path)
            at.run(timeout=30)
            return at
        finally:
            if os.path.exists(app_path):
                os.remove(app_path)
            if os.path.exists(script_path):
                os.remove(script_path)

    def test_fall9_vorhandenes_sheet_wird_importiert(self):
        at = self._run_tab_mit_bytes(self._build_workbook_mit_sheet())
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))
        texts = [m.value for m in at.markdown] + [t.value for t in at.text]
        self.assertIn("ANZAHL_TERMINE=1", texts)
        successes = [s.body for s in at.success]
        self.assertTrue(any("aus dem vorhandenen Tabellenblatt" in s for s in successes))

    def test_fall10_fehlendes_sheet_kein_fehler_leere_liste(self):
        at = self._run_tab_mit_bytes(self._build_workbook_ohne_sheet())
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))
        texts = [m.value for m in at.markdown] + [t.value for t in at.text]
        self.assertIn("ANZAHL_TERMINE=0", texts)


class TestTerminHinzufuegenUndBearbeitenUeberEchteWidgets(unittest.TestCase):
    """Regressionstest über die ECHTEN Streamlit-Widgets (nicht nur die
    zugrundeliegende Logik) für "Termin hinzufügen" und "Termin bearbeiten" -
    deckt insbesondere ab, dass das Zurücksetzen der Formularfelder nach dem
    Absenden nicht erneut zu einem StreamlitWidgetAlreadyInstantiatedError
    führt (bereits einmal in dieser Form aufgetreten und behoben, siehe
    app.py -> _render_termin_hinzufuegen_formular)."""

    def setUp(self):
        self.path = os.path.join(os.path.dirname(__file__), "_tmp_widget_flow.py")
        with open(self.path, "w", encoding="utf-8") as f:
            f.write(
                "import app as app_module\n"
                "config = app_module.load_config()\n"
                "app_module.render_routenplanung_tab(config, None)\n"
            )

    def tearDown(self):
        if os.path.exists(self.path):
            os.remove(self.path)

    def test_mehrfaches_hinzufuegen_und_bearbeiten_ohne_absturz(self):
        at = AppTest.from_file(self.path)
        at.run(timeout=30)

        def text_inputs():
            return {ti.label: ti for ti in at.text_input}

        def date_inputs():
            return {di.label: di for di in at.date_input}

        def time_inputs():
            return {ti.label: ti for ti in at.time_input}

        def buttons():
            return {b.label: b for b in at.button}

        # Erster Termin hinzufügen (bewusst ohne Straße, siehe Fall 4).
        text_inputs()["Unternehmensname"].set_value("Hotel Eins").run()
        text_inputs()["Straße"].set_value("").run()
        text_inputs()["PLZ"].set_value("88045").run()
        text_inputs()["Ort"].set_value("Friedrichshafen").run()
        date_inputs()["Termin_Datum"].set_value(dt.date(2026, 10, 15)).run()
        time_inputs()["Termin_Uhrzeit"].set_value(dt.time(9, 0)).run()
        buttons()["Termin hinzufügen"].click().run(timeout=30)
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))

        # Zweiter Termin direkt danach - das Formular muss korrekt
        # zurückgesetzt worden sein (keine WidgetAlreadyInstantiated-Fehler).
        text_inputs()["Unternehmensname"].set_value("Hotel Zwei").run()
        text_inputs()["Straße"].set_value("Seestr. 2").run()
        text_inputs()["PLZ"].set_value("88045").run()
        text_inputs()["Ort"].set_value("Friedrichshafen").run()
        date_inputs()["Termin_Datum"].set_value(dt.date(2026, 10, 15)).run()
        time_inputs()["Termin_Uhrzeit"].set_value(dt.time(11, 0)).run()
        buttons()["Termin hinzufügen"].click().run(timeout=30)
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))

        warnings = [w.body for w in at.warning]
        self.assertTrue(any("Hotel Eins: Straße fehlt." in w for w in warnings))

        # "Hotel Eins" bearbeiten und die fehlende Straße ergänzen (Fall 4).
        edit_buttons = [b for b in at.button if b.key and b.key.startswith("edit_")]
        self.assertEqual(len(edit_buttons), 2)
        edit_buttons[0].click().run(timeout=30)
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))

        tid = edit_buttons[0].key.split("edit_", 1)[1]
        edit_strasse = next(ti for ti in at.text_input if ti.key == f"bearb_strasse_{tid}")
        edit_strasse.set_value("Nachträglich ergänzt").run(timeout=30)
        fertig_button = next(b for b in at.button if b.key == f"bearb_fertig_{tid}")
        fertig_button.click().run(timeout=30)
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))

        warnings_danach = [w.body for w in at.warning]
        self.assertFalse(any("Hotel Eins: Straße fehlt." in w for w in warnings_danach))

        # Einen der beiden Termine löschen (Fall 7).
        delete_buttons = [b for b in at.button if b.key and b.key.startswith("delete_")]
        self.assertEqual(len(delete_buttons), 2)
        delete_buttons[0].click().run(timeout=30)
        self.assertEqual(len(at.exception), 0, msg=list(at.exception))
        self.assertEqual(len([b for b in at.button if b.key and b.key.startswith("delete_")]), 1)


if __name__ == "__main__":
    unittest.main()
